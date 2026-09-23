import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 Widescreen standard
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color Palette
    PRIMARY_DARK = RGBColor(26, 54, 93)       # Deep Navy #1A365D
    PRIMARY_BLUE = RGBColor(37, 99, 235)      # Royal Blue #2563EB
    PRIMARY_LIGHT = RGBColor(239, 246, 255)   # Soft Ice Blue #EFF6FF
    ACCENT_TEAL = RGBColor(13, 148, 136)      # Teal #0D9488
    ACCENT_ORANGE = RGBColor(234, 88, 12)     # Orange #EA580C
    TEXT_DARK = RGBColor(30, 41, 59)          # Slate 800 #1E293B
    TEXT_MUTED = RGBColor(100, 116, 139)      # Slate 500 #64748B
    CARD_BG = RGBColor(255, 255, 255)         # White
    CARD_BORDER = RGBColor(226, 232, 240)     # Gray 200 #E2E8F0
    ACCENT_GREEN = RGBColor(22, 163, 74)      # Green 600

    def add_header(slide, title_text, category_text="TIN HỌC 10 - BÀI 7: THỰC HÀNH SỬ DỤNG THIẾT BỊ SỐ THÔNG DỤNG"):
        # Header banner container
        top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
        top_bar.fill.solid()
        top_bar.fill.fore_color.rgb = PRIMARY_DARK
        top_bar.line.color.rgb = PRIMARY_DARK

        tf = top_bar.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_top = Inches(0.12)
        tf.margin_right = Inches(0.3)
        tf.margin_bottom = Inches(0.1)

        p_cat = tf.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.name = "Calibri"
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = RGBColor(147, 197, 253)

        p_title = tf.add_paragraph()
        p_title.text = title_text
        p_title.font.name = "Calibri"
        p_title.font.size = Pt(21)
        p_title.font.bold = True
        p_title.font.color.rgb = RGBColor(255, 255, 255)

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=CARD_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)
        return card

    # ==========================================
    # SLIDE 1: TIÊU ĐỀ
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = PRIMARY_DARK
    bg1.line.fill.background()

    card_title = add_card(slide1, Inches(1.2), Inches(1.1), Inches(10.933), Inches(5.3), RGBColor(255, 255, 255), PRIMARY_BLUE)
    tf1 = card_title.text_frame
    tf1.word_wrap = True
    tf1.margin_left = Inches(0.8)
    tf1.margin_right = Inches(0.8)
    tf1.margin_top = Inches(0.6)

    p0 = tf1.paragraphs[0]
    p0.alignment = PP_ALIGN.CENTER
    p0.text = "TIN HỌC 10 • BỘ SÁCH KẾT NỐI TRI THỨC VỚI CUỘC SỐNG • CHỦ ĐỀ 1"
    p0.font.name = "Calibri"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = PRIMARY_BLUE

    p1 = tf1.add_paragraph()
    p1.alignment = PP_ALIGN.CENTER
    p1.text = "BÀI 7: THỰC HÀNH SỬ DỤNG\nTHIẾT BỊ SỐ THÔNG DỤNG"
    p1.font.name = "Calibri"
    p1.font.size = Pt(32)
    p1.font.bold = True
    p1.font.color.rgb = PRIMARY_DARK

    p_line = tf1.add_paragraph()
    p_line.alignment = PP_ALIGN.CENTER
    p_line.text = "━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━"
    p_line.font.size = Pt(11)
    p_line.font.color.rgb = ACCENT_TEAL

    p2 = tf1.add_paragraph()
    p2.alignment = PP_ALIGN.CENTER
    p2.text = "Nội dung toàn diện: Lý thuyết PDA & 4 Nhiệm vụ thực hành • HĐ1 • LT1, LT2 • VD1, VD2, VD3"
    p2.font.name = "Calibri"
    p2.font.size = Pt(14)
    p2.font.bold = True
    p2.font.color.rgb = ACCENT_ORANGE

    p3 = tf1.add_paragraph()
    p3.alignment = PP_ALIGN.CENTER
    p3.text = "\nBài trình chiếu phục vụ học tập và thuyết trình trên lớp"
    p3.font.name = "Calibri"
    p3.font.size = Pt(13)
    p3.font.color.rgb = TEXT_MUTED

    # ==========================================
    # SLIDE 2: MỤC TIÊU & CẤU TRÚC BÀI HỌC
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_header(slide2, "MỤC TIÊU BÀI HỌC VÀ LỘ TRÌNH THỰC HÀNH")

    col_w = Inches(3.7)
    gap = Inches(0.3)
    start_x = Inches(0.8)
    top_y = Inches(1.8)
    h = Inches(5.1)

    c1 = add_card(slide2, start_x, top_y, col_w, h, PRIMARY_LIGHT, PRIMARY_BLUE)
    tf_c1 = c1.text_frame
    tf_c1.word_wrap = True
    tf_c1.margin_left = Inches(0.25)
    tf_c1.margin_right = Inches(0.25)
    tf_c1.margin_top = Inches(0.3)
    p = tf_c1.paragraphs[0]
    p.text = "🎯 MỤC TIÊU CẦN ĐẠT"
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_DARK

    bullets_c1 = [
        "Biết được các thiết bị số cá nhân (PDA) thông dụng gồm những gì.",
        "Nắm bắt các tính năng tiêu biểu của thiết bị số di động.",
        "Khai thác và sử dụng thành thạo ứng dụng, dữ liệu trên smartphone, tablet.",
        "Thành thạo cấu trúc thư mục tệp tin (DCIM, Camera) và thao tác quản lý tệp.",
        "Kết nối chia sẻ dữ liệu máy tính và lưu trữ đám mây an toàn."
    ]
    for b in bullets_c1:
        p = tf_c1.add_paragraph()
        p.text = "• " + b
        p.font.name = "Calibri"
        p.font.size = Pt(12.5)
        p.font.color.rgb = TEXT_DARK

    c2 = add_card(slide2, start_x + col_w + gap, top_y, col_w, h, CARD_BG, ACCENT_TEAL)
    tf_c2 = c2.text_frame
    tf_c2.word_wrap = True
    tf_c2.margin_left = Inches(0.25)
    tf_c2.margin_right = Inches(0.25)
    tf_c2.margin_top = Inches(0.3)
    p = tf_c2.paragraphs[0]
    p.text = "📚 NỘI DUNG LÝ THUYẾT & THỰC HÀNH"
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_DARK

    bullets_c2 = [
        "1. Trợ thủ số cá nhân (PDA): Khái niệm, vai trò & HĐ1.",
        "2. Nhiệm vụ 1: Nhận biết nút bấm vật lý & khởi động máy.",
        "3. Nhiệm vụ 2: Màn hình chính, thanh trạng thái, điều hướng.",
        "4. Nhiệm vụ 3: Khám phá ứng dụng thiết yếu & cài đặt thêm.",
        "5. Nhiệm vụ 4: Ứng dụng quản lý tệp tin và cây thư mục.",
        "6. Luyện tập (LT1, LT2) & Vận dụng (VD1, VD2, VD3)."
    ]
    for b in bullets_c2:
        p = tf_c2.add_paragraph()
        p.text = "✔ " + b
        p.font.name = "Calibri"
        p.font.size = Pt(12.5)
        p.font.color.rgb = TEXT_DARK

    c3 = add_card(slide2, start_x + (col_w + gap) * 2, top_y, col_w, h, CARD_BG, ACCENT_ORANGE)
    tf_c3 = c3.text_frame
    tf_c3.word_wrap = True
    tf_c3.margin_left = Inches(0.25)
    tf_c3.margin_right = Inches(0.25)
    tf_c3.margin_top = Inches(0.3)
    p = tf_c3.paragraphs[0]
    p.text = "⚙️ THIẾT BỊ & CÔNG CỤ"
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_DARK

    bullets_c3 = [
        "Điện thoại thông minh (chạy Android hoặc iOS/iPhone).",
        "Máy tính để bàn / Laptop có cổng USB.",
        "Cáp kết nối dữ liệu USB (Type-C / Lightning).",
        "Kết nối mạng Internet / Wi-Fi trường học hoặc gia đình.",
        "Tài khoản đám mây (Google Drive hoặc OneDrive)."
    ]
    for b in bullets_c3:
        p = tf_c3.add_paragraph()
        p.text = "➤ " + b
        p.font.name = "Calibri"
        p.font.size = Pt(12.5)
        p.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 3: MỤC 1 - TRỢ THỦ SỐ CÁ NHÂN & HĐ1
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_header(slide3, "1. TRỢ THỦ SỐ CÁ NHÂN (PDA) & HOẠT ĐỘNG 1 (HĐ1)")

    c_l3 = add_card(slide3, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1), CARD_BG, PRIMARY_BLUE)
    tf_l3 = c_l3.text_frame
    tf_l3.word_wrap = True
    tf_l3.margin_left = Inches(0.3)
    tf_l3.margin_right = Inches(0.3)
    tf_l3.margin_top = Inches(0.3)

    p = tf_l3.paragraphs[0]
    p.text = "A. KHÁI NIỆM TRỢ THỦ SỐ CÁ NHÂN (PDA)"
    p.font.name = "Calibri"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE

    p = tf_l3.add_paragraph()
    p.text = "• PDA (Personal Digital Assistant) là các thiết bị số tích hợp nhiều chức năng và phần mềm ứng dụng hữu ích phục vụ người dùng trong đời sống hàng ngày.\n• Đặc điểm then chốt:\n  - Kích thước nhỏ gọn, tính di động cao, dễ mang theo người.\n  - Chạy hệ điều hành riêng biệt.\n  - Khả năng kết nối mạng không dây (Wi-Fi, 4G/5G, Bluetooth)."
    p.font.size = Pt(12.5)
    p.font.color.rgb = TEXT_DARK

    p = tf_l3.add_paragraph()
    p.text = "\nB. CÁC THIẾT BỊ PDA THÔNG DỤNG HIỆN NAY:"
    p.font.name = "Calibri"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_DARK

    devices = [
        "Điện thoại thông minh (Smartphone): Phổ biến nhất, đa năng nhất.",
        "Máy tính bảng (Tablet): Màn hình lớn, tối ưu làm việc, học tập, vẽ.",
        "Đồng hồ thông minh (Smartwatch): Theo dõi sức khỏe, thông báo tiện ích.",
        "Máy đọc sách (E-reader): Màn hình E-Ink chuyên dụng đọc sách báo."
    ]
    for d in devices:
        p = tf_l3.add_paragraph()
        p.text = "• " + d
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK

    c_r3 = add_card(slide3, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1), PRIMARY_LIGHT, ACCENT_TEAL)
    tf_r3 = c_r3.text_frame
    tf_r3.word_wrap = True
    tf_r3.margin_left = Inches(0.3)
    tf_r3.margin_right = Inches(0.3)
    tf_r3.margin_top = Inches(0.3)

    p = tf_r3.paragraphs[0]
    p.text = "C. HOẠT ĐỘNG 1: TÌM HIỂU THIẾT BỊ & ỨNG DỤNG"
    p.font.name = "Calibri"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_TEAL

    p = tf_r3.add_paragraph()
    p.text = "❓ Câu hỏi HĐ1 (SGK trang 33): Hãy liệt kê một số thiết bị trợ thủ số cá nhân và các ứng dụng tiêu biểu đi kèm.\n\nTrả lời chi tiết:\n1. Ứng dụng ban đầu (Thời kỳ đầu của PDA):\n   - Sổ lịch công tác, đồng hồ để xem thời gian, đặt giờ báo thức.\n   - Sổ danh bạ ghi địa chỉ, số điện thoại liên hệ.\n   - Danh sách việc cần làm (to-do list), sổ ghi nhớ (notes), máy tính bỏ túi.\n\n2. Ứng dụng hiện đại ngày nay (Tích hợp siêu đa nhiệm):\n   - Nghe nhạc, ghi âm, xem phim, chụp ảnh, quay phim sắc nét.\n   - Gọi điện, nhắn tin đa phương tiện, định vị tìm đường (GPS).\n   - Điều khiển thiết bị điện tử từ xa, thanh toán thẻ / mã QR."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 4: HĐ1 - HỆ ĐIỀU HÀNH & CÂU HỎI CỦNG CỐ
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_header(slide4, "HĐ1: HỆ ĐIỀU HÀNH PHỔ BIẾN & CÂU HỎI CỦNG CỐ")

    c_l4 = add_card(slide4, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1), CARD_BG, PRIMARY_BLUE)
    tf_l4 = c_l4.text_frame
    tf_l4.word_wrap = True
    tf_l4.margin_left = Inches(0.3)
    tf_l4.margin_right = Inches(0.3)
    tf_l4.margin_top = Inches(0.3)

    p = tf_l4.paragraphs[0]
    p.text = "📱 HAI HỆ ĐIỀU HÀNH THỐNG TRỊ PDA"
    p.font.name = "Calibri"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE

    os_info = [
        "1. Android (Hãng Google):",
        "   - Hệ điều hành mã nguồn mở phổ biến nhất thế giới.",
        "   - Trang bị trên hầu hết điện thoại của Samsung, Xiaomi, Oppo, Vivo...",
        "   - Kho ứng dụng khổng lồ: Google Play Store.",
        "2. iOS / iPadOS (Hãng Apple):",
        "   - Hệ điều hành độc quyền chạy riêng trên iPhone và iPad.",
        "   - Tối ưu hóa mượt mà, bảo mật cao, hệ sinh thái đồng bộ hoàn hảo.",
        "   - Kho ứng dụng: Apple App Store.",
        "3. Cổng giao tiếp & kết nối trên PDA:",
        "   - Cổng sạc và truyền tệp: USB Type-C, Lightning.",
        "   - Kết nối không dây: Wi-Fi, Bluetooth, 4G / 5G, NFC."
    ]
    for line in os_info:
        p = tf_l4.add_paragraph()
        p.text = line
        p.font.size = Pt(12)
        p.font.bold = line.startswith("1.") or line.startswith("2.") or line.startswith("3.")
        p.font.color.rgb = PRIMARY_DARK if p.font.bold else TEXT_DARK

    c_r4 = add_card(slide4, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1), PRIMARY_LIGHT, ACCENT_ORANGE)
    tf_r4 = c_r4.text_frame
    tf_r4.word_wrap = True
    tf_r4.margin_left = Inches(0.3)
    tf_r4.margin_right = Inches(0.3)
    tf_r4.margin_top = Inches(0.3)

    p = tf_r4.paragraphs[0]
    p.text = "❓ CÂU HỎI CỦNG CỐ HĐ1 (SGK TRANG 34)"
    p.font.name = "Calibri"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_ORANGE

    p = tf_r4.add_paragraph()
    p.text = "Kết nối nào KHÔNG phải là kết nối phổ biến trên các PDA hiện nay?"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_DARK

    opts = ["A. Wi-Fi", "B. Bluetooth", "C. Hồng ngoại", "D. USB"]
    for o in opts:
        p = tf_r4.add_paragraph()
        p.text = "   " + o
        p.font.size = Pt(14)
        p.font.bold = o.startswith("C.")
        p.font.color.rgb = ACCENT_GREEN if o.startswith("C.") else TEXT_DARK

    p = tf_r4.add_paragraph()
    p.text = "\n💡 ĐÁP ÁN: C. HỒNG NGOẠI"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    p = tf_r4.add_paragraph()
    p.text = "• Giải thích: Cổng hồng ngoại có tốc độ rất chậm, cự ly ngắn (< 1m) và phải hướng thẳng vào nhau. Hiện nay, Wi-Fi và Bluetooth với tốc độ cao, không cần ngắm thẳng đã thay thế hoàn toàn hồng ngoại."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 5: NHIỆM VỤ 1 - NÚT BẤM VẬT LÝ & KHỞI ĐỘNG
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_header(slide5, "2. THỰC HÀNH: NHIỆM VỤ 1 - NÚT BẤM VẬT LÝ & KHỞI ĐỘNG")

    c_nv1_l = add_card(slide5, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1), CARD_BG, PRIMARY_BLUE)
    tf_n1l = c_nv1_l.text_frame
    tf_n1l.word_wrap = True
    tf_n1l.margin_left = Inches(0.3)
    tf_n1l.margin_right = Inches(0.3)
    tf_n1l.margin_top = Inches(0.3)

    p = tf_n1l.paragraphs[0]
    p.text = "🔘 CÁC NÚT BẤM VẬT LÝ TIÊU BIỂU"
    p.font.name = "Calibri"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE

    buttons = [
        "1. Nút khóa / Nguồn (Power / Lock Key):",
        "   - Nhấn 1 lần: Tắt hoặc bật sáng màn hình để khóa/mở khóa.",
        "   - Nhấn giữ lâu: Hiện menu Tắt nguồn (Power off), Khởi động lại (Restart), hoặc gọi khẩn cấp SOS.",
        "2. Cụm nút điều chỉnh âm lượng (Volume +/-):",
        "   - Tăng hoặc giảm âm lượng cuộc gọi, chuông báo, âm thanh đa phương tiện.",
        "   - Kết hợp nút Nguồn để chụp ảnh màn hình (Screenshot).",
        "3. Phím công tắc gạt âm thanh (Gạt rung trên iPhone):",
        "   - Chuyển đổi siêu nhanh giữa chế độ Chuông và Rung/Im lặng."
    ]
    for b in buttons:
        p = tf_n1l.add_paragraph()
        p.text = b
        p.font.size = Pt(12)
        p.font.bold = b.startswith("1.") or b.startswith("2.") or b.startswith("3.")
        p.font.color.rgb = PRIMARY_DARK if p.font.bold else TEXT_DARK

    c_nv1_r = add_card(slide5, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1), PRIMARY_LIGHT, ACCENT_TEAL)
    tf_n1r = c_nv1_r.text_frame
    tf_n1r.word_wrap = True
    tf_n1r.margin_left = Inches(0.3)
    tf_n1r.margin_right = Inches(0.3)
    tf_n1r.margin_top = Inches(0.3)

    p = tf_n1r.paragraphs[0]
    p.text = "🚀 CÁC BƯỚC THỰC HIỆN NHIỆM VỤ 1"
    p.font.name = "Calibri"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_TEAL

    steps_nv1 = [
        "• Bước 1: Quan sát vị trí bố trí các nút bấm ở cạnh trái, cạnh phải thân máy điện thoại đang sử dụng.",
        "• Bước 2: Bấm nút khóa để khởi động/đánh thức màn hình điện thoại.",
        "• Bước 3: Mở khóa bằng mã PIN, mật khẩu, hình vẽ (pattern) hoặc sinh trắc học (vân tay, nhận diện khuôn mặt).",
        "• Bước 4: Kiểm tra hệ điều hành đang dùng:",
        "   - Vào Cài đặt (Settings) ➔ Thông tin điện thoại (About phone) ➔ Thông tin phần mềm (Software info) để xem phiên bản Android hoặc iOS."
    ]
    for s in steps_nv1:
        p = tf_n1r.add_paragraph()
        p.text = s
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 6: NHIỆM VỤ 2 - MÀN HÌNH LÀM VIỆC & ĐIỀU HƯỚNG
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_header(slide6, "2. THỰC HÀNH: NHIỆM VỤ 2 - MÀN HÌNH CHÍNH & ĐIỀU HƯỚNG")

    c_nv2_l = add_card(slide6, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1), CARD_BG, PRIMARY_BLUE)
    tf_n2l = c_nv2_l.text_frame
    tf_n2l.word_wrap = True
    tf_n2l.margin_left = Inches(0.3)
    tf_n2l.margin_right = Inches(0.3)
    tf_n2l.margin_top = Inches(0.3)

    p = tf_n2l.paragraphs[0]
    p.text = "🖥️ CẤU TRÚC MÀN HÌNH CHÍNH (HOME SCREEN)"
    p.font.name = "Calibri"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE

    zones = [
        "1. Thanh trạng thái (Status Bar - Ở trên cùng):",
        "   - Hiển thị thời gian hiện tại, vạch sóng di động, cột sóng Wi-Fi/4G, tỉ lệ phần trăm pin còn lại, biểu tượng thông báo ứng dụng.",
        "2. Khu vực lưới biểu tượng (App Icons Grid):",
        "   - Nơi bố trí các biểu tượng ứng dụng và tiện ích (widgets: đồng hồ, thời tiết). Vuốt sang trái/phải để chuyển trang.",
        "3. Thanh truy cập nhanh (Dock Bar - Dưới màn hình):",
        "   - Chứa các app quan trọng nhất (Điện thoại, Tin nhắn, Danh bạ, Trình duyệt) cố định ở mọi trang màn hình."
    ]
    for z in zones:
        p = tf_n2l.add_paragraph()
        p.text = z
        p.font.size = Pt(12)
        p.font.bold = z.startswith("1.") or z.startswith("2.") or z.startswith("3.")
        p.font.color.rgb = PRIMARY_DARK if p.font.bold else TEXT_DARK

    c_nv2_r = add_card(slide6, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1), PRIMARY_LIGHT, ACCENT_TEAL)
    tf_n2r = c_nv2_r.text_frame
    tf_n2r.word_wrap = True
    tf_n2r.margin_left = Inches(0.3)
    tf_n2r.margin_right = Inches(0.3)
    tf_n2r.margin_top = Inches(0.3)

    p = tf_n2r.paragraphs[0]
    p.text = "🧭 THANH ĐIỀU HƯỚNG & CỬ CHỈ CẢM ỨNG"
    p.font.name = "Calibri"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_TEAL

    navs = [
        "1. Ba phím điều hướng ảo (Navigation Bar trên Android):",
        "   - Nút Home (Giữa): Bấm để quay về ngay màn hình chính.",
        "   - Nút Quay lại / Back (Trái hoặc Phải): Trở về trang/màn hình thao tác trước đó.",
        "   - Nút Tổng quan / Overview (Ba gạch hoặc ô vuông): Hiện danh sách các ứng dụng đang chạy nền để chuyển đổi hoặc đóng app.",
        "2. Điều hướng cử chỉ hiện đại (Gesture Navigation):",
        "   - Vuốt từ cạnh dưới lên: Về Home.",
        "   - Vuốt từ mép trái/phải vào trong: Quay lại (Back).",
        "   - Vuốt từ dưới lên và giữ 1 giây: Mở đa nhiệm ứng dụng."
    ]
    for n in navs:
        p = tf_n2r.add_paragraph()
        p.text = n
        p.font.size = Pt(12)
        p.font.bold = n.startswith("1.") or n.startswith("2.")
        p.font.color.rgb = PRIMARY_DARK if p.font.bold else TEXT_DARK

    # ==========================================
    # SLIDE 7: NHIỆM VỤ 3 - BIỂU TƯỢNG VÀ ỨNG DỤNG
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_header(slide7, "2. THỰC HÀNH: NHIỆM VỤ 3 - CÁC NHÓM ỨNG DỤNG TRÊN ĐIỆN THOẠI")

    box_w = Inches(3.7)
    gap = Inches(0.3)
    start_x = Inches(0.8)
    top_y = Inches(1.8)
    h = Inches(5.1)

    c_g1 = add_card(slide7, start_x, top_y, box_w, h, CARD_BG, PRIMARY_BLUE)
    tf_g1 = c_g1.text_frame
    tf_g1.word_wrap = True
    tf_g1.margin_left = Inches(0.25)
    tf_g1.margin_right = Inches(0.25)
    tf_g1.margin_top = Inches(0.3)
    p = tf_g1.paragraphs[0]
    p.text = "📞 ỨNG DỤNG THIẾT YẾU"
    p.font.name = "Calibri"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE

    g1_items = [
        "Đây là các ứng dụng nền tảng phục vụ chức năng nghe - gọi cơ bản:",
        "• Điện thoại (Phone / Call): Quay số, thực hiện cuộc gọi, xem nhật ký liên lạc.",
        "• Nhắn tin (Messages / SMS): Soạn và nhận tin nhắn văn bản, tin nhắn đa phương tiện.",
        "• Danh bạ (Contacts): Quản lý lưu trữ tên, số điện thoại, email, địa chỉ của người thân, bạn bè."
    ]
    for it in g1_items:
        p = tf_g1.add_paragraph()
        p.text = it
        p.font.size = Pt(12.5)
        p.font.color.rgb = TEXT_DARK

    c_g2 = add_card(slide7, start_x + box_w + gap, top_y, box_w, h, PRIMARY_LIGHT, ACCENT_TEAL)
    tf_g2 = c_g2.text_frame
    tf_g2.word_wrap = True
    tf_g2.margin_left = Inches(0.25)
    tf_g2.margin_right = Inches(0.25)
    tf_g2.margin_top = Inches(0.3)
    p = tf_g2.paragraphs[0]
    p.text = "📷 ỨNG DỤNG MẶC ĐỊNH SẴN CÓ"
    p.font.name = "Calibri"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_TEAL

    g2_items = [
        "Nhà sản xuất cài đặt sẵn nhằm hỗ trợ cuộc sống hàng ngày:",
        "• Máy ảnh (Camera) & Bộ sưu tập ảnh (Photos/Gallery): Ghi lại và quản lý album ảnh, video.",
        "• Trình duyệt Web (Chrome, Safari): Truy cập tra cứu thông tin.",
        "• Ứng dụng Email: Nhận gửi thư điện tử học tập, công việc.",
        "• Tiện ích công cụ: Đồng hồ báo thức, Lịch, Máy tính số học, Ghi chú."
    ]
    for it in g2_items:
        p = tf_g2.add_paragraph()
        p.text = it
        p.font.size = Pt(12.5)
        p.font.color.rgb = TEXT_DARK

    c_g3 = add_card(slide7, start_x + (box_w + gap) * 2, top_y, box_w, h, CARD_BG, ACCENT_ORANGE)
    tf_g3 = c_g3.text_frame
    tf_g3.word_wrap = True
    tf_g3.margin_left = Inches(0.25)
    tf_g3.margin_right = Inches(0.25)
    tf_g3.margin_top = Inches(0.3)
    p = tf_g3.paragraphs[0]
    p.text = "🌐 ỨNG DỤNG CÀI ĐẶT THÊM"
    p.font.name = "Calibri"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_ORANGE

    g3_items = [
        "Tải về từ Chợ ứng dụng (Google Play Store hoặc App Store):",
        "• Học tập trực tuyến: Zoom, MS Teams, Google Meet, Google Classroom.",
        "• Mạng xã hội & Liên lạc: Zalo, Facebook, YouTube, Messenger.",
        "• Lưu trữ đám mây: Google Drive, Microsoft OneDrive, Dropbox.",
        "• Bản đồ & Giao thông: Google Maps, ứng dụng đặt xe công nghệ."
    ]
    for it in g3_items:
        p = tf_g3.add_paragraph()
        p.text = it
        p.font.size = Pt(12.5)
        p.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 8: NHIỆM VỤ 4 - HỆ THỐNG QUẢN LÝ TỆP TIN
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_header(slide8, "2. THỰC HÀNH: NHIỆM VỤ 4 - QUẢN LÝ TỆP TRÊN DI ĐỘNG")

    c_nv4_l = add_card(slide8, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1), CARD_BG, PRIMARY_BLUE)
    tf_n4l = c_nv4_l.text_frame
    tf_n4l.word_wrap = True
    tf_n4l.margin_left = Inches(0.3)
    tf_n4l.margin_right = Inches(0.3)
    tf_n4l.margin_top = Inches(0.3)

    p = tf_n4l.paragraphs[0]
    p.text = "📂 CẤU TRÚC PHÂN CẤP TỆP TIN DI ĐỘNG"
    p.font.name = "Calibri"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE

    p = tf_n4l.add_paragraph()
    p.text = "• Hệ thống tệp trên điện thoại thông minh tổ chức theo cấu trúc hình cây phân cấp hoàn toàn tương tự máy tính:\n  - Thư mục gốc (Root) ➔ Các thư mục con ➔ Các tệp tin dữ liệu.\n• Các không gian lưu trữ chính:\n  - Bộ nhớ trong (Internal Storage): Bộ nhớ tích hợp trên máy.\n  - Thẻ nhớ ngoài (SD Card): Nếu điện thoại hỗ trợ khe cắm thẻ.\n  - Ổ đĩa đám mây (Cloud): Tích hợp trực tiếp trong ứng dụng tệp (như Google Drive, OneDrive, iCloud)."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_DARK

    p = tf_n4l.add_paragraph()
    p.text = "\n📸 Thư mục lưu trữ ảnh chuẩn:"
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = ACCENT_ORANGE

    p = tf_n4l.add_paragraph()
    p.text = "• Toàn bộ ảnh chụp từ camera mặc định được lưu tại thư mục: DCIM (viết tắt của Digital Camera Images) ➔ Thư mục con Camera."
    p.font.size = Pt(12)
    p.font.color.rgb = TEXT_DARK

    c_nv4_r = add_card(slide8, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1), PRIMARY_LIGHT, ACCENT_TEAL)
    tf_n4r = c_nv4_r.text_frame
    tf_n4r.word_wrap = True
    tf_n4r.margin_left = Inches(0.3)
    tf_n4r.margin_right = Inches(0.3)
    tf_n4r.margin_top = Inches(0.3)

    p = tf_n4r.paragraphs[0]
    p.text = "🛠️ CÁC THAO TÁC QUẢN TRỊ TỆP CƠ BẢN"
    p.font.name = "Calibri"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_TEAL

    ops = [
        "1. Mở xem tệp:",
        "   - Nhấp chạm vào tệp để xem nội dung toàn màn hình.",
        "2. Chọn nhiều tệp (Multi-select):",
        "   - Chạm và giữ lâu (Long-press) vào 1 tệp, sau đó tích chọn các tệp còn lại.",
        "3. Sao chép (Copy) & Di chuyển (Move):",
        "   - Chọn tệp ➔ Nhấn Sao chép hoặc Di chuyển ➔ Điều hướng tới thư mục đích ➔ Nhấn 'Dán ở đây' (Paste).",
        "4. Đổi tên (Rename) & Chia sẻ (Share):",
        "   - Đổi tên tệp giúp quản lý dễ dàng; chia sẻ qua Bluetooth, Zalo, Drive.",
        "5. Xóa tệp (Delete):",
        "   - Chọn tệp ➔ Nhấn biểu tượng Thùng rác để giải phóng dung lượng."
    ]
    for op in ops:
        p = tf_n4r.add_paragraph()
        p.text = op
        p.font.size = Pt(11.5)
        p.font.bold = op.startswith("1.") or op.startswith("2.") or op.startswith("3.") or op.startswith("4.") or op.startswith("5.")
        p.font.color.rgb = PRIMARY_DARK if p.font.bold else TEXT_DARK

    # ==========================================
    # SLIDE 9: LUYỆN TẬP 1 (LT1) - PHÂN BIỆT SMARTPHONE & ĐIỆN THOẠI THƯỜNG
    # ==========================================
    slide9 = prs.slides.add_slide(blank_layout)
    add_header(slide9, "LUYỆN TẬP 1 (LT1): PHÂN BIỆT SMARTPHONE VÀ ĐIỆN THOẠI THƯỜNG")

    c_lt1 = add_card(slide9, Inches(1.2), Inches(1.8), Inches(10.933), Inches(5.1), CARD_BG, PRIMARY_BLUE)
    tf_lt1 = c_lt1.text_frame
    tf_lt1.word_wrap = True
    tf_lt1.margin_left = Inches(0.4)
    tf_lt1.margin_right = Inches(0.4)
    tf_lt1.margin_top = Inches(0.3)

    p = tf_lt1.paragraphs[0]
    p.text = "❓ CÂU HỎI TRẮC NGHIỆM (SGK TRANG 37):"
    p.font.name = "Calibri"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_ORANGE

    p = tf_lt1.add_paragraph()
    p.text = "Điện thoại thông minh khác với điện thoại di động thường ở điểm nào?"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_DARK

    lt1_opts = [
        "A. Có danh bạ.",
        "B. Có thể nhắn tin.",
        "C. Có thể kết nối Internet.",
        "D. Có hệ điều hành và có thể chạy được một số ứng dụng."
    ]
    for o in lt1_opts:
        p = tf_lt1.add_paragraph()
        p.text = "   " + o
        p.font.size = Pt(14)
        p.font.bold = (o.startswith("D."))
        p.font.color.rgb = ACCENT_GREEN if o.startswith("D.") else TEXT_DARK

    p = tf_lt1.add_paragraph()
    p.text = "\n✅ ĐÁP ÁN ĐÚNG: D"
    p.font.size = Pt(17)
    p.font.bold = True
    p.font.color.rgb = ACCENT_GREEN

    p = tf_lt1.add_paragraph()
    p.text = "🔍 GIẢI THÍCH CHI TIẾT BẢN CHẤT CÔNG NGHỆ:"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_DARK

    lt1_detail = [
        "• Điện thoại thường (Feature Phone / Cục gạch): Chạy phần mềm nhúng cố định (firmware). Người dùng không thể cài thêm phần mềm mới bên ngoài. Dù máy vẫn có danh bạ, nhắn tin SMS và thậm chí một số máy duyệt được web 2G cơ bản.",
        "• Điện thoại thông minh (Smartphone): Có Hệ điều hành riêng (Android, iOS) với nhân hệ thống mạnh mẽ, hỗ trợ chạy đa nhiệm, và sở hữu kho ứng dụng (App Store, Google Play) cho phép người dùng tùy ý cài đặt hàng triệu ứng dụng theo nhu cầu."
    ]
    for d in lt1_detail:
        p = tf_lt1.add_paragraph()
        p.text = d
        p.font.size = Pt(12.5)
        p.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 10: LUYỆN TẬP 2 (LT2) - CHỤP ẢNH & THAO TÁC QUẢN LÝ TỆP
    # ==========================================
    slide10 = prs.slides.add_slide(blank_layout)
    add_header(slide10, "LUYỆN TẬP 2 (LT2): THỰC HÀNH CHỤP ẢNH, DUYỆT TỆP & XÓA ẢNH")

    box_w = Inches(2.1)
    box_gap = Inches(0.25)
    box_top = Inches(1.8)
    box_h = Inches(5.1)

    lt2_steps = [
        ("BƯỚC 1", "Chụp ảnh mới", "• Mở ứng dụng Máy ảnh (Camera) trên máy.\n• Nhấn chụp một bức ảnh trong lớp học.\n• Tệp ảnh tự động được lưu vào bộ nhớ trong."),
        ("BƯỚC 2", "Mở Quản lý tệp", "• Tìm ứng dụng quản lý tệp trên máy.\n• Tên ứng dụng: File Manager, My Files (Samsung) hoặc Tệp (Files trên iPhone).\n• Mở ứng dụng lên."),
        ("BƯỚC 3", "Định vị DCIM", "• Chọn: Bộ nhớ trong (Internal Storage).\n• Tìm và mở thư mục DCIM (Digital Camera Images).\n• Nhấp chọn thư mục con Camera."),
        ("BƯỚC 4", "Mở xem ảnh", "• Tìm bức ảnh vừa chụp (thường đứng đầu danh sách, đuôi .jpg hoặc .heic).\n• Chạm vào tệp để xem toàn màn hình và kiểm tra thông tin."),
        ("BƯỚC 5", "Xóa tệp ảnh", "• Chạm và giữ lâu vào tệp ảnh (Long press).\n• Chọn biểu tượng Thùng rác (Xóa / Delete).\n• Bấm 'Xác nhận xóa' để hoàn tất bài thực hành.")
    ]

    for i, (st, title, desc) in enumerate(lt2_steps):
        x = Inches(0.8) + (box_w + box_gap) * i
        bg = PRIMARY_LIGHT if i % 2 == 0 else CARD_BG
        border = PRIMARY_BLUE if i == 0 or i == 4 else ACCENT_TEAL
        c_step = add_card(slide10, x, box_top, box_w, box_h, bg, border)
        tf_s = c_step.text_frame
        tf_s.word_wrap = True
        tf_s.margin_left = Inches(0.15)
        tf_s.margin_right = Inches(0.15)
        tf_s.margin_top = Inches(0.2)

        p = tf_s.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        p.text = st
        p.font.name = "Calibri"
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.color.rgb = ACCENT_ORANGE

        p = tf_s.add_paragraph()
        p.alignment = PP_ALIGN.CENTER
        p.text = title
        p.font.name = "Calibri"
        p.font.size = Pt(15)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_DARK

        p = tf_s.add_paragraph()
        p.text = "─────────"
        p.font.size = Pt(8)
        p.font.color.rgb = TEXT_MUTED

        p = tf_s.add_paragraph()
        p.text = desc
        p.font.name = "Calibri"
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 11: VẬN DỤNG 1 (VD1) - KẾT NỐI VỚI MÁY TÍNH
    # ==========================================
    slide11 = prs.slides.add_slide(blank_layout)
    add_header(slide11, "VẬN DỤNG 1 (VD1): KẾT NỐI ĐIỆN THOẠI VỚI MÁY TÍNH ĐỂ SAO CHÉP ẢNH")

    c_vd1_l = add_card(slide11, Inches(0.8), Inches(1.8), Inches(6.8), Inches(5.1), CARD_BG, PRIMARY_BLUE)
    tf_vl = c_vd1_l.text_frame
    tf_vl.word_wrap = True
    tf_vl.margin_left = Inches(0.3)
    tf_vl.margin_right = Inches(0.3)
    tf_vl.margin_top = Inches(0.3)

    p = tf_vl.paragraphs[0]
    p.text = "📋 QUY TRÌNH THỰC HIỆN QUA CÁP USB:"
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_DARK

    vd1_steps = [
        "Bước 1 (Cắm cáp): Dùng cáp USB, cắm một đầu vào điện thoại, đầu kia cắm vào cổng USB trên máy tính.",
        "Bước 2 (Cấp quyền MTP): Vuốt màn hình điện thoại xuống, chọn chế độ 'Truyền tệp / Truyền tệp tin' (File Transfer / MTP). (Trên iPhone: Nhấn 'Tin cậy máy tính này').",
        "Bước 3 (Mở trên máy tính): Mở This PC (File Explorer) trên máy tính, nhấp đúp vào biểu tượng tên điện thoại.",
        "Bước 4 (Định vị ảnh): Vào: Bộ nhớ trong (Internal Storage) ➔ DCIM ➔ Camera.",
        "Bước 5 (Sao chép): Chọn các ảnh cần chuyển ➔ Chuột phải chọn Copy (Ctrl+C) ➔ Chuyển sang thư mục đích trên máy tính ➔ Chuột phải chọn Paste (Ctrl+V)."
    ]
    for s in vd1_steps:
        p = tf_vl.add_paragraph()
        p.text = s
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK

    c_vd1_r = add_card(slide11, Inches(7.9), Inches(1.8), Inches(4.6), Inches(5.1), PRIMARY_LIGHT, ACCENT_ORANGE)
    tf_vr = c_vd1_r.text_frame
    tf_vr.word_wrap = True
    tf_vr.margin_left = Inches(0.3)
    tf_vr.margin_right = Inches(0.3)
    tf_vr.margin_top = Inches(0.3)

    p = tf_vr.paragraphs[0]
    p.text = "⚠️ LƯU Ý KỸ THUẬT & MỞ RỘNG:"
    p.font.name = "Calibri"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_ORANGE

    notes = [
        "Dùng cáp truyền dữ liệu: Cáp sạc đơn thuần không truyền được tệp.",
        "Mở khóa màn hình: Điện thoại phải mở khóa thì máy tính mới đọc được tệp tin.",
        "Truyền tệp không dây (Mở rộng): Có thể dùng Quick Share (Android sang Windows), AirDrop (Apple), hoặc gửi file qua Zalo / Telegram Web."
    ]
    for n in notes:
        p = tf_vr.add_paragraph()
        p.text = "• " + n
        p.font.size = Pt(12.5)
        p.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 12: VẬN DỤNG 2 (VD2) - LƯU TRỮ ĐÁM MÂY (CLOUD)
    # ==========================================
    slide12 = prs.slides.add_slide(blank_layout)
    add_header(slide12, "VẬN DỤNG 2 (VD2): LƯU TRỮ ẢNH TRÊN DỊCH VỤ ĐÁM MÂY (GOOGLE DRIVE / ONEDRIVE)")

    c_vd2_l = add_card(slide12, Inches(0.8), Inches(1.8), Inches(6.8), Inches(5.1), CARD_BG, ACCENT_TEAL)
    tf_v2l = c_vd2_l.text_frame
    tf_v2l.word_wrap = True
    tf_v2l.margin_left = Inches(0.3)
    tf_v2l.margin_right = Inches(0.3)
    tf_v2l.margin_top = Inches(0.3)

    p = tf_v2l.paragraphs[0]
    p.text = "☁️ 5 BƯỚC THỰC HIỆN LƯU TRỮ ĐÁM MÂY:"
    p.font.name = "Calibri"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_DARK

    v2_steps = [
        "Bước 1: Mở ứng dụng Google Drive trên điện thoại (hoặc vào drive.google.com trên máy tính).",
        "Bước 2: Đăng nhập bằng tài khoản Google (Gmail) hoặc tài khoản Microsoft.",
        "Bước 3: Nhấn nút dấu cộng '+' ➔ Chọn 'Thư mục mới' ➔ Đặt tên thư mục: 'Anh_Thuc_Hanh_Tin_10'.",
        "Bước 4: Mở thư mục vừa tạo ➔ Nhấn dấu '+' ➔ Chọn 'Tải lên' (Upload) ➔ Chọn Hình ảnh ➔ Tích chọn những bức ảnh muốn sao lưu.",
        "Bước 5: Chờ tải lên hoàn tất ➔ Dữ liệu đã được bảo vệ vĩnh viễn trên đám mây an toàn."
    ]
    for s in v2_steps:
        p = tf_v2l.add_paragraph()
        p.text = s
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK

    c_vd2_r = add_card(slide12, Inches(7.9), Inches(1.8), Inches(4.6), Inches(5.1), PRIMARY_LIGHT, PRIMARY_BLUE)
    tf_v2r = c_vd2_r.text_frame
    tf_v2r.word_wrap = True
    tf_v2r.margin_left = Inches(0.3)
    tf_v2r.margin_right = Inches(0.3)
    tf_v2r.margin_top = Inches(0.3)

    p = tf_v2r.paragraphs[0]
    p.text = "🌟 LỢI ÍCH CỦA LƯU TRỮ ĐÁM MÂY:"
    p.font.name = "Calibri"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE

    benefits = [
        "Truy cập mọi lúc, mọi nơi: Xem và tải ảnh trên mọi thiết bị kết nối mạng.",
        "An toàn tuyệt đối: Không sợ mất dữ liệu khi hỏng hóc hoặc mất điện thoại.",
        "Giải phóng bộ nhớ: Cho phép xóa ảnh trên máy sau khi đã sao lưu đám mây.",
        "Chia sẻ siêu tốc: Tạo liên kết (link) gửi ảnh cho thầy cô và bạn bè chỉ với 1 chạm."
    ]
    for b in benefits:
        p = tf_v2r.add_paragraph()
        p.text = "✔ " + b
        p.font.size = Pt(12.5)
        p.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 13: VẬN DỤNG 3 (VD3) & TỔNG KẾT BÀI HỌC
    # ==========================================
    slide13 = prs.slides.add_slide(blank_layout)
    add_header(slide13, "VẬN DỤNG 3 (VD3) & TỔNG KẾT KIẾN THỨC CỐT LÕI BÀI 7")

    c_vd3 = add_card(slide13, Inches(0.8), Inches(1.8), Inches(5.7), Inches(5.1), CARD_BG, ACCENT_ORANGE)
    tf_v3 = c_vd3.text_frame
    tf_v3.word_wrap = True
    tf_v3.margin_left = Inches(0.3)
    tf_v3.margin_right = Inches(0.3)
    tf_v3.margin_top = Inches(0.3)

    p = tf_v3.paragraphs[0]
    p.text = "💬 VẬN DỤNG 3 (SGK TRANG 37)"
    p.font.name = "Calibri"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = ACCENT_ORANGE

    p = tf_v3.add_paragraph()
    p.text = "Đề bài: Gửi ảnh qua phần mềm hỗ trợ học tập trực tuyến (Zoom / Meet / Teams / Zalo)."
    p.font.size = Pt(13)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_DARK

    vd3_guide = [
        "• Trên Zoom / Google Meet / MS Teams:",
        "  - Trong phòng học trực tuyến, mở khung Trò chuyện (Chat).",
        "  - Nhấn vào biểu tượng đính kèm tệp (hình kẹp ghim 📎) hoặc biểu tượng ảnh.",
        "  - Chọn ảnh chụp bài thực hành từ thư viện điện thoại/máy tính ➔ Nhấn Gửi.",
        "• Trên Zalo / Nhóm học tập:",
        "  - Nhấp vào biểu tượng Hình ảnh ➔ Chọn chế độ HD để giữ nguyên độ nét ➔ Gửi cho giáo viên bộ môn."
    ]
    for g in vd3_guide:
        p = tf_v3.add_paragraph()
        p.text = g
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK

    c_sum13 = add_card(slide13, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.1), PRIMARY_LIGHT, PRIMARY_DARK)
    tf_s13 = c_sum13.text_frame
    tf_s13.word_wrap = True
    tf_s13.margin_left = Inches(0.3)
    tf_s13.margin_right = Inches(0.3)
    tf_s13.margin_top = Inches(0.3)

    p = tf_s13.paragraphs[0]
    p.text = "📌 4 ĐIỀU GHI NHỚ CỐT LÕI CỦA BÀI 7"
    p.font.name = "Calibri"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_DARK

    keynotes = [
        "1. Trợ thủ số (PDA): Nhỏ gọn, cơ động, hệ điều hành riêng và tích hợp kết nối mạng không dây tốc độ cao.",
        "2. Điểm phân biệt Smartphone: Có hệ điều hành (Android, iOS) và cài đặt được kho ứng dụng phong phú.",
        "3. Quản lý tệp chuẩn xác: Cấu trúc thư mục hình cây; ảnh chụp lưu tại 'Internal Storage / DCIM / Camera'.",
        "4. Kỹ năng dữ liệu thời đại số: Thành thạo cả truyền tệp có dây (cáp USB - chế độ MTP) và sao lưu đám mây không dây (Google Drive, OneDrive)."
    ]
    for k in keynotes:
        p = tf_s13.add_paragraph()
        p.text = k
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_DARK

    # ==========================================
    # SLIDE 14: KẾT THÚC & Q&A
    # ==========================================
    slide14 = prs.slides.add_slide(blank_layout)
    bg14 = slide14.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg14.fill.solid()
    bg14.fill.fore_color.rgb = PRIMARY_DARK
    bg14.line.fill.background()

    card_end = add_card(slide14, Inches(1.8), Inches(1.3), Inches(9.733), Inches(4.9), RGBColor(255, 255, 255), PRIMARY_BLUE)
    tf_end = card_end.text_frame
    tf_end.word_wrap = True
    tf_end.margin_left = Inches(0.5)
    tf_end.margin_right = Inches(0.5)
    tf_end.margin_top = Inches(0.7)

    p = tf_end.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    p.text = "CẢM ƠN THẦY CÔ VÀ CÁC BẠN\nĐÃ CHÚ Ý THEO DÕI BÀI THUYẾT TRÌNH!"
    p.font.name = "Calibri"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_DARK

    p = tf_end.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.text = "━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━━"
    p.font.size = Pt(12)
    p.font.color.rgb = ACCENT_TEAL

    p = tf_end.add_paragraph()
    p.alignment = PP_ALIGN.CENTER
    p.text = "RẤT MONG NHẬN ĐƯỢC NHỮNG Ý KIẾN ĐÓNG GÓP VÀ CÂU HỎI TỪ MỌI NGƯỜI!\n\n❓ PHẦN HỎI ĐÁP (Q & A)"
    p.font.name = "Calibri"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = PRIMARY_BLUE

    # Save
    output_path = "/home/truonglangquan/idk-lop10/idk-tin/Bai_7_Tin_Hoc_10_Thuc_Hanh_Thiet_Bi_So.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully with {len(prs.slides)} slides at: {output_path}")

if __name__ == "__main__":
    create_presentation()
